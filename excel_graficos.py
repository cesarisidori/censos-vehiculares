"""
excel_graficos.py
-----------------
Gráficos nativos (openpyxl) para el Excel de salida del censo.

El motor (`procesar_censo3.py`) no conoce openpyxl: escribe hojas de datos
auxiliares y luego llama a un callback (`al_escribir`) que recibe el libro
abierto. Este módulo construye ese callback.

`preparar_graficos()` devuelve:

- `hojas_extra` : dict {nombre_hoja: DataFrame} con los datos en formato
  "ancho" (período + una columna por serie), listos para que los escriba el
  motor.
- `al_escribir` : callable(workbook) que agrega los gráficos nativos
  (editables en Excel / OnlyOffice / LibreOffice) en una hoja "Gráficos".

Funciones puras: no dependen de Streamlit ni del motor.
"""

import pandas as pd

from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint, Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from graficos import PALETA_DEFAULT

# Geometría de la grilla de gráficos (en la hoja "Gráficos").
_COLUMNAS_ANCLA = (2, 11)   # B y K
_FILA_INICIO = 4
_PASO_FILA = 18
_ANCHO_CM = 13
_ALTO_CM = 8


# ─────────────────────────────────────────────────────────────
# Preparación de datos (formato ancho)
# ─────────────────────────────────────────────────────────────
def _orden(df, col_periodo, periodos):
    """Períodos en orden (el provisto o los únicos ordenados)."""
    if periodos is not None:
        return list(periodos)
    return sorted(df[col_periodo].dropna().unique())


def _ancho_por_sentido(df, col_periodo, periodos):
    """Período + una columna por sentido (volumen sumado entre calles)."""
    datos = {"Periodo": _orden(df, col_periodo, periodos)}
    for sentido in sorted(df["sentido"].unique()):
        serie = (df[df["sentido"] == sentido]
                 .groupby(col_periodo)["volumen"].sum())
        datos[str(sentido)] = [int(serie.get(p, 0)) for p in datos["Periodo"]]
    return pd.DataFrame(datos)


def _ancho_por_calle(df, col_periodo, periodos):
    """Período + una columna por calle/sentido + 'Calle | Ambos'."""
    datos = {"Periodo": _orden(df, col_periodo, periodos)}
    for calle in sorted(df["calle"].unique()):
        sub = df[df["calle"] == calle]
        for sentido in sorted(sub["sentido"].unique()):
            serie = (sub[sub["sentido"] == sentido]
                     .groupby(col_periodo)["volumen"].sum())
            datos[f"{calle} | {sentido}"] = [
                int(serie.get(p, 0)) for p in datos["Periodo"]]
        total = sub.groupby(col_periodo)["volumen"].sum()
        datos[f"{calle} | Ambos"] = [
            int(total.get(p, 0)) for p in datos["Periodo"]]
    return pd.DataFrame(datos)


def _tabla_composicion(serie_hp, serie_total):
    """Tipos (fila) con el volumen en hora pico y en el censo completo."""
    tipos = []
    for serie in (serie_hp, serie_total):
        if serie is None:
            continue
        for tipo in serie.index:
            if tipo not in tipos:
                tipos.append(tipo)

    def _valores(serie):
        if serie is None:
            return [0] * len(tipos)
        return [int(serie.get(tipo, 0)) for tipo in tipos]

    tabla = pd.DataFrame({
        "Tipo": tipos,
        "Hora pico": _valores(serie_hp),
        "Censo completo": _valores(serie_total),
    })
    tabla = tabla[(tabla["Hora pico"] > 0) | (tabla["Censo completo"] > 0)]
    return tabla.reset_index(drop=True)


def _tabla_pesados(pesados_hp):
    """Movimiento + '% pesados' (solo los mayores a cero, ordenado)."""
    if pesados_hp is None or "% pesados" not in pesados_hp.index:
        return None
    pct = pesados_hp.loc["% pesados"]
    pct = pct[[c for c in pct.index if c != "Total"]]
    pct = pd.to_numeric(pct, errors="coerce").dropna()
    pct = pct[pct > 0].sort_values(ascending=True)
    if pct.empty:
        return None
    return pd.DataFrame({
        "Movimiento": [str(m) for m in pct.index],
        "% pesados": [float(v) for v in pct.values],
    })


def _spec(tipo, hoja, titulo, cat, series):
    """Especificación mínima de un gráfico."""
    return {"tipo": tipo, "hoja": hoja, "titulo": titulo,
            "cat": cat, "series": list(series)}


# ─────────────────────────────────────────────────────────────
# Construcción del callback
# ─────────────────────────────────────────────────────────────
def preparar_graficos(df_cs, agg_hora=None, periodos_hp=None,
                      serie_hp=None, serie_total=None, pesados_hp=None,
                      titulo="", subtitulo="", paleta=None):
    """
    Devuelve (hojas_extra, al_escribir) para `exportar`/`exportar_buffer`.

    - `df_cs`        : DataFrame largo (hora_inicio, calle, sentido, volumen).
    - `agg_hora`     : DataFrame agregado por hora entera (o None).
    - `periodos_hp`  : lista de bloques de 15 min dentro de la hora pico.
    - `serie_hp`     : Series (tipo -> volumen) en hora pico (o None).
    - `serie_total`  : Series (tipo -> volumen) del censo completo (o None).
    - `pesados_hp`   : tabla del motor con fila "% pesados" (o None).
    - `titulo`       : título base del censo.
    """
    paleta = list(paleta) if paleta else list(PALETA_DEFAULT)
    hojas, specs = {}, []

    # 1) Líneas: distribución por 15 min, censo completo.
    if df_cs is not None and not df_cs.empty:
        ancho = _ancho_por_sentido(df_cs, "hora_inicio", None)
        hojas["Graf_15min"] = ancho
        specs.append(_spec(
            "linea", "Graf_15min",
            "Distribución por 15 minutos — censo completo",
            "Periodo", list(ancho.columns[1:])))

    # 2) Barras: volumen por hora, una por calle.
    if agg_hora is not None and not agg_hora.empty:
        ancho = _ancho_por_calle(agg_hora, "hora_entera", None)
        hojas["Graf_Hora"] = ancho
        for calle in sorted(agg_hora["calle"].unique()):
            series = [c for c in ancho.columns if c.startswith(f"{calle} | ")]
            specs.append(_spec("barra", "Graf_Hora",
                               f"Volumen por hora — {calle}",
                               "Periodo", series))

    # 3) Barras: 15 min dentro de la hora pico, una por calle.
    if df_cs is not None and not df_cs.empty and periodos_hp:
        df_hp = df_cs[df_cs["hora_inicio"].isin(list(periodos_hp))]
        if not df_hp.empty:
            ancho = _ancho_por_calle(df_hp, "hora_inicio", list(periodos_hp))
            hojas["Graf_15min_HP"] = ancho
            for calle in sorted(df_hp["calle"].unique()):
                series = [c for c in ancho.columns
                          if c.startswith(f"{calle} | ")]
                specs.append(_spec(
                    "barra", "Graf_15min_HP",
                    f"Volumen cada 15 min (hora pico) — {calle}",
                    "Periodo", series))

    # 4) Tortas: composición por tipo (hora pico y censo completo).
    comp = _tabla_composicion(serie_hp, serie_total)
    if not comp.empty:
        hojas["Graf_Composicion"] = comp
        if int(comp["Hora pico"].sum()) > 0:
            specs.append(_spec("torta", "Graf_Composicion",
                               "Composición por tipo — hora pico",
                               "Tipo", ["Hora pico"]))
        if int(comp["Censo completo"].sum()) > 0:
            specs.append(_spec("torta", "Graf_Composicion",
                               "Composición por tipo — censo completo",
                               "Tipo", ["Censo completo"]))

    # 5) Barras horizontales: % de pesados por movimiento.
    pes = _tabla_pesados(pesados_hp)
    if pes is not None and not pes.empty:
        hojas["Graf_Pesados"] = pes
        specs.append(_spec("barra_h", "Graf_Pesados",
                           "% de pesados por movimiento — hora pico",
                           "Movimiento", ["% pesados"]))

    def al_escribir(workbook):
        _crear_graficos(workbook, specs, paleta, titulo, subtitulo)

    return hojas, al_escribir


# ─────────────────────────────────────────────────────────────
# Creación de los gráficos en el libro
# ─────────────────────────────────────────────────────────────
def _columna(encabezados, nombre):
    """Índice (1-based) de la columna con ese encabezado, o None."""
    return encabezados.get(nombre)


def _hex(color):
    """Color válido para OOXML (sin '#')."""
    return str(color).lstrip("#")


def _pintar_series(chart, paleta):
    """Colorea las series (y los marcadores de línea) con la paleta."""
    for i, serie in enumerate(chart.series):
        color = _hex(paleta[i % len(paleta)])
        props = GraphicalProperties(solidFill=color)
        serie.graphicalProperties = props
        if isinstance(chart, LineChart):
            props.line = LineProperties(solidFill=color, w=20000)
            serie.marker = Marker(symbol="circle", size=6)


def _pintar_torta(chart, paleta, n_categorias):
    """Colorea cada porción de la torta/dona."""
    dolientes = []
    for i in range(n_categorias):
        punto = DataPoint(idx=i)
        punto.graphicalProperties = GraphicalProperties(
            solidFill=_hex(paleta[i % len(paleta)]))
        dolientes.append(punto)
    if chart.series:
        chart.series[0].data_points = dolientes


def _nueva_barra(horizontal):
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.grouping = "clustered"
    chart.style = 10
    chart.legend.position = "b"
    return chart


def _nueva_linea():
    chart = LineChart()
    chart.style = 10
    chart.legend.position = "b"
    return chart


def _nueva_torta():
    chart = DoughnutChart()
    chart.style = 10
    chart.legend.position = "r"
    chart.holeSize = 40
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showPercent = True
    chart.dataLabels.showVal = False
    chart.dataLabels.showCatName = False
    chart.dataLabels.showSerName = False
    chart.dataLabels.showLegendKey = False
    chart.dataLabels.showBubbleSize = False
    return chart


def _crear_graficos(workbook, specs, paleta, titulo="", subtitulo=""):
    """Agrega todos los gráficos en una hoja 'Gráficos'."""
    if not specs:
        return

    hoja = workbook.create_sheet("Gráficos")
    if titulo:
        hoja["A1"] = titulo
        hoja["A1"].font = Font(size=14, bold=True)
    if subtitulo:
        hoja["A2"] = subtitulo
        hoja["A2"].font = Font(size=11, italic=True)

    for i, spec in enumerate(specs):
        fuente = workbook[spec["hoja"]]
        if fuente.max_row < 2:
            continue

        encabezados = {c.value: c.column for c in fuente[1]}
        cat_col = _columna(encabezados, spec["cat"])
        cols = [c for c in (_columna(encabezados, s)
                            for s in spec["series"]) if c]
        if not cat_col or not cols:
            continue

        data = Reference(fuente, min_col=min(cols), max_col=max(cols),
                         min_row=1, max_row=fuente.max_row)
        cats = Reference(fuente, min_col=cat_col,
                         min_row=2, max_row=fuente.max_row)

        if spec["tipo"] == "linea":
            chart = _nueva_linea()
        elif spec["tipo"] == "torta":
            chart = _nueva_torta()
        else:
            chart = _nueva_barra(spec["tipo"] == "barra_h")

        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        # Las categorías son texto (horas, calles, tipos): se referencian
        # como strRef para que Excel/OnlyOffice las traten como etiquetas.
        letra = get_column_letter(cat_col)
        ref_cat = (f"'{spec['hoja']}'!${letra}$2:${letra}${fuente.max_row}")
        for serie in chart.series:
            serie.cat = AxDataSource(strRef=StrRef(f=ref_cat))
        chart.title = spec["titulo"]

        if spec["tipo"] == "torta":
            _pintar_torta(chart, paleta, fuente.max_row - 1)
        else:
            _pintar_series(chart, paleta)
            chart.x_axis.delete = False
            chart.y_axis.delete = False

        chart.width = _ANCHO_CM
        chart.height = _ALTO_CM
        col = _COLUMNAS_ANCLA[i % len(_COLUMNAS_ANCLA)]
        fila = _FILA_INICIO + (i // len(_COLUMNAS_ANCLA)) * _PASO_FILA
        hoja.add_chart(chart, f"{get_column_letter(col)}{fila}")
