"""
generar_plantilla.py
-------------------
Genera una plantilla Excel de censo vehicular completamente parametrizable.

USO:
    python generar_plantilla.py config.json
    python generar_plantilla.py            # usa config por defecto (ejemplo)
"""

import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ═══════════════════════════════════════════════════════════════
# CONFIGURACIÓN POR DEFECTO (podés sobreescribir con un JSON)
# ═══════════════════════════════════════════════════════════════
CONFIG_DEFAULT = {
    "archivo_salida": "Plantilla_Censos.xlsx",
    "metadatos": {
        "fecha": "2026-10-09",
        "dia": "viernes",
        "ubicacion": "RP 18 y El Caldén",
        "turno": "Mañana",
        "lugar": "Las Carolinas",
    },
    "hora_inicio": "07:30",
    "duracion_horas": 2,
    "tipos_vehiculo": ["Autos", "Bicicletas", "Veh. Pesados", "TUP"],
    # Cada movimiento: calle, sentido, nro_mov
    "movimientos": [
        {"calle": "RP 18",     "sentido": "Sur-Norte",  "nro_mov": 4},
        {"calle": "RP 18",     "sentido": "Sur-Norte",  "nro_mov": 1},
        {"calle": "RP 18",     "sentido": "Sur-Norte",  "nro_mov": 14},
        {"calle": "RP 18",     "sentido": "Norte-Sur",  "nro_mov": 13},
        {"calle": "RP 18",     "sentido": "Norte-Sur",  "nro_mov": 2},
        {"calle": "RP 18",     "sentido": "Norte-Sur",  "nro_mov": 3},
        {"calle": "EL CALDEN", "sentido": "Este-Oeste", "nro_mov": 12},
        {"calle": "EL CALDEN", "sentido": "Este-Oeste", "nro_mov": 15},
        {"calle": "EL CALDEN", "sentido": "Este-Oeste", "nro_mov": 11},
        {"calle": "EL CALDEN", "sentido": "Oeste-Este", "nro_mov": 6},
        {"calle": "EL CALDEN", "sentido": "Oeste-Este", "nro_mov": 16},
        {"calle": "EL CALDEN", "sentido": "Oeste-Este", "nro_mov": 5},
    ],
    # Tipos que NO se cuentan para hora pico (se excluyen del total)
    "tipos_excluidos_hora_pico": ["Bicicletas", "Peatones"],
}

# ═══════════════════════════════════════════════════════════════
# ESTILOS
# ═══════════════════════════════════════════════════════════════
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FILL_CALLE   = PatternFill("solid", fgColor="D9E1F2")
FILL_SENTIDO = PatternFill("solid", fgColor="E2EFDA")
FILL_HORA    = PatternFill("solid", fgColor="FFF2CC")
FILL_TIPO    = PatternFill("solid", fgColor="FCE4D6")

# ═══════════════════════════════════════════════════════════════
# UTILIDADES
# ═══════════════════════════════════════════════════════════════
def sumar_minutos(hhmm, minutos):
    h, m = map(int, hhmm.split(":"))
    total = h * 60 + m + minutos
    return f"{(total // 60) % 24:02d}:{total % 60:02d}"

def generar_bloques_15min(hora_inicio, duracion_horas):
    return [
        (sumar_minutos(hora_inicio, i * 15), sumar_minutos(hora_inicio, (i + 1) * 15))
        for i in range(duracion_horas * 4)
    ]

# ═══════════════════════════════════════════════════════════════
# HOJA CONFIG (se escribe tal cual, para que el parser la lea)
# ═══════════════════════════════════════════════════════════════
def hoja_config(wb, cfg):
    ws = wb.create_sheet("CONFIG")
    ws["A1"] = "PARÁMETRO"; ws["B1"] = "VALOR"
    for c in "AB":
        ws[f"{c}1"].font = BOLD
        ws[f"{c}1"].fill = FILL_CALLE

    r = 2
    # Metadatos
    for k, v in cfg["metadatos"].items():
        ws[f"A{r}"] = k; ws[f"B{r}"] = v; r += 1
    # Parámetros de tiempo
    ws[f"A{r}"] = "hora_inicio"; ws[f"B{r}"] = cfg["hora_inicio"]; r += 1
    ws[f"A{r}"] = "duracion_horas"; ws[f"B{r}"] = cfg["duracion_horas"]; r += 1

    # Tipos de vehículo
    r += 1
    ws[f"A{r}"] = "TIPOS DE VEHÍCULO"; ws[f"A{r}"].font = BOLD; r += 1
    for t in cfg["tipos_vehiculo"]:
        ws[f"A{r}"] = t; r += 1

    # Tipos excluidos del cálculo
    r += 1
    ws[f"A{r}"] = "TIPOS EXCLUIDOS HORA PICO"; ws[f"A{r}"].font = BOLD; r += 1
    for t in cfg["tipos_excluidos_hora_pico"]:
        ws[f"A{r}"] = t; r += 1

    # Movimientos
    r += 1
    ws[f"A{r}"] = "MOVIMIENTOS"; ws[f"A{r}"].font = BOLD; r += 1
    ws[f"A{r}"] = "Calle"; ws[f"B{r}"] = "Sentido"; ws[f"C{r}"] = "nro_mov"
    for c in "ABC": ws[f"{c}{r}"].font = BOLD
    r += 1
    for m in cfg["movimientos"]:
        ws[f"A{r}"] = m["calle"]
        ws[f"B{r}"] = m["sentido"]
        ws[f"C{r}"] = m["nro_mov"]
        r += 1

    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 12

# ═══════════════════════════════════════════════════════════════
# HOJA CENSO (formulario de carga)
# ═══════════════════════════════════════════════════════════════
def hoja_censo(wb, cfg):
    ws = wb.create_sheet("CENSO", 0)
    movs = cfg["movimientos"]
    tipos = cfg["tipos_vehiculo"]
    n_mov = len(movs)

    # Encabezados en 3 filas
    ws.cell(1, 1, "Hora Inicio"); ws.cell(1, 2, "Hora Fin"); ws.cell(1, 3, "Tipo")
    for c in (1, 2, 3):
        ws.cell(1, c).font = BOLD
        ws.cell(1, c).fill = FILL_HORA
        ws.cell(1, c).alignment = CENTER
        ws.cell(1, c).border = BORDER

    for j, m in enumerate(movs):
        col = 4 + j
        ws.cell(1, col, m["calle"]).fill = FILL_CALLE
        ws.cell(2, col, m["sentido"]).fill = FILL_SENTIDO
        ws.cell(3, col, m["nro_mov"])
        for fila in (1, 2, 3):
            c = ws.cell(fila, col)
            c.alignment = CENTER
            c.border = BORDER
            c.font = BOLD if fila == 1 else Font()

    # Datos
    bloques = generar_bloques_15min(cfg["hora_inicio"], cfg["duracion_horas"])
    r = 4
    for hora_ini, hora_fin in bloques:
        for tipo in tipos:
            ws.cell(r, 1, hora_ini)
            ws.cell(r, 2, hora_fin)
            ws.cell(r, 3, tipo).fill = FILL_TIPO
            for c in (1, 2, 3):
                ws.cell(r, c).border = BORDER
                ws.cell(r, c).alignment = CENTER
            for j in range(n_mov):
                cell = ws.cell(r, 4 + j, 0)
                cell.border = BORDER
                cell.alignment = CENTER
            r += 1

    # Anchos y freeze
    ws.column_dimensions["A"].width = 11
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 14
    for j in range(n_mov):
        ws.column_dimensions[get_column_letter(4 + j)].width = 9
    ws.freeze_panes = "D4"

# ═══════════════════════════════════════════════════════════════
# HOJA INSTRUCCIONES
# ═══════════════════════════════════════════════════════════════
def hoja_instrucciones(wb, cfg):
    ws = wb.create_sheet("INSTRUCCIONES")
    lineas = [
        "CÓMO USAR ESTA PLANTILLA",
        "",
        "1) Hoja CONFIG: NO tocar salvo que cambie el censo.",
        "   Ahí se definen calles, sentidos, movimientos, tipos y duración.",
        "",
        "2) Hoja CENSO: acá se cargan los conteos.",
        "   - Cada fila = 15 min × tipo de vehículo.",
        "   - Cada columna (D en adelante) = un movimiento.",
        "   - Los encabezados indican Calle / Sentido / nro_mov.",
        "   - NO agregar filas Total: el script las calcula.",
        "   - NO cambiar el orden de columnas.",
        "",
        "3) Guardar como .xlsx y pasarlo a procesar_censo.py.",
        "",
        f"Duración configurada: {cfg['duracion_horas']} h desde las {cfg['hora_inicio']}",
        f"Tipos de vehículo: {', '.join(cfg['tipos_vehiculo'])}",
        f"Movimientos: {len(cfg['movimientos'])}",
    ]
    for i, txt in enumerate(lineas, start=1):
        ws.cell(i, 1, txt)
    ws.column_dimensions["A"].width = 80

# ═══════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════
def main():
    cfg = CONFIG_DEFAULT
    if len(sys.argv) > 1:
        with open(sys.argv[1], "r", encoding="utf-8") as f:
            cfg = json.load(f)

    wb = Workbook()
    wb.remove(wb.active)
    hoja_config(wb, cfg)
    hoja_censo(wb, cfg)
    hoja_instrucciones(wb, cfg)

    wb.save(cfg["archivo_salida"])
    print(f"✅ Plantilla generada: {cfg['archivo_salida']}")

if __name__ == "__main__":
    main()