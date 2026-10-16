"""
reportes.py
-----------
Generación de reportes descargables del censo:

- PDF  (reportlab)   → `generar_pdf(ctx)`
- DOCX (python-docx) → `generar_docx(ctx)` — se abre en Word y en LibreOffice

Las figuras Plotly se convierten a PNG con mirage (`fig_a_png`). Los
croquis (SVG) se rasterizan a PNG con svg2png-py/resvg (`svg_a_png`).
Ambos se embeben como imágenes en el PDF y el DOCX.

`ctx` es un dict con:
    titulo, subtitulo                 str
    fecha (date), dia                 date / str
    ubicacion, lugar, analista        str
    fhp_data                          dict (ver procesar_censo3.calcular_fhp)
    resumen_agr, resumen_plano,       DataFrames (o None)
    pesados_hp, totales
    excluidos                         list[str]
    imagenes                          list[(titulo_seccion, png_bytes | None)]
    croquis                           list[(titulo_seccion, svg_str)]

Tipografía:
- PDF  → DejaVu Sans TTF **embebida** (fonts/ del proyecto → sistema →
  Arial de Windows → fallback Helvetica con acentos vía WinAnsiEncoding).
- DOCX → Arial en todos los estilos base (portable Windows/Linux).
"""

import io
import re
import subprocess
import sys
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image as PILImage

# Emojis y símbolos decorativos: en el PDF se eliminan para que el
# texto quede siempre limpio (ninguna de las fuentes candidatas los trae).
_EMOJI = re.compile(
    "[\U0001F000-\U0001FAFF\U00002190-\U000024FF\uFE0F\u200D]+"
)

# Candidatas a fuente embebida, en orden de prioridad:
# (regular, bold) — primero las del proyecto, luego sistema, luego Windows.
_FUENTES = [
    (Path(__file__).parent / "fonts" / "DejaVuSans.ttf",
     Path(__file__).parent / "fonts" / "DejaVuSans-Bold.ttf"),
    (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
     Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")),
    (Path("C:/Windows/Fonts/arial.ttf"),
     Path("C:/Windows/Fonts/arialbd.ttf")),
]


# ─────────────────────────────────────────────────────────────
# Helpers comunes
# ─────────────────────────────────────────────────────────────
def _sin_emoji(texto):
    """Quita emojis y normaliza a str."""
    return _EMOJI.sub("", str(texto or "")).strip()


def _fmt(valor):
    """Formatea un valor de celda de forma compacta."""
    if isinstance(valor, bool):
        return str(valor)
    if isinstance(valor, float):
        return f"{valor:g}"
    return str(valor)


def _titulo_plano(texto):
    """
    Aplana un título de Plotly a UNA sola línea, para que mirage no
    superponga el título con el subtítulo (`<br><sub>…</sub>`).

    'TÍTULO<br><sub>SUBTÍTULO</sub>' → 'TÍTULO — SUBTÍTULO'
    """
    if not texto:
        return texto
    plano = re.sub(r"<br\s*/?>\s*<sub>", " — ", str(texto), flags=re.I)
    plano = re.sub(r"<br\s*/?>", " — ", plano, flags=re.I)
    plano = re.sub(r"<[^>]+>", "", plano)
    return plano.strip()


def fig_a_png(fig, ancho=None, escala=2, timeout=10):
    """
    Convierte una figura Plotly a PNG con mirage.
    Devuelve bytes o None si mirage no está disponible o falla.

    - Pasa SIEMPRE `width` y `height` para que mirage no use sus valores
      por defecto (que deforman el aspect ratio de la figura).
    - Si se indica `ancho`, el alto se recalcula manteniendo la
      proporción original de `fig.layout`.
    - El título se aplana sobre una COPIA (el `fig` original no se toca,
      así en pantalla sigue con 2 líneas).
    """
    try:
        # Mirage requiere que la figura sea serializable (sin ndarray de NumPy)
        import json
        import plotly.graph_objects as go
        import plotly.io as pio
        import mirage

        # Convertir figura a JSON puro para evitar problemas con ndarray
        fig_json = pio.to_json(fig)
        fig_serializable = go.Figure(json.loads(fig_json))

        # Título a UNA línea sobre la copia (el original no se modifica)
        fig_serializable.update_layout(
            title_text=_titulo_plano(fig_serializable.layout.title.text),
            title_font=dict(size=12))

        # Dimensiones: derivadas de la figura, con proporción resuelta
        w0 = fig_serializable.layout.width or 700
        h0 = fig_serializable.layout.height or 450
        ancho_final = int(ancho) if ancho else int(w0)
        alto_final = max(1, round(ancho_final * h0 / w0))

        kwargs = {"format": "png", "scale": escala,
                  "width": ancho_final, "height": alto_final}

        # Usar mirage para convertir a PNG
        png_bytes = mirage.to_image(fig_serializable, **kwargs)
        return png_bytes
    except Exception:
        return None


def _png_dims(png):
    """(ancho_px, alto_px) de un PNG en bytes."""
    with PILImage.open(io.BytesIO(png)) as im:
        return im.size


# ─────────────────────────────────────────────────────────────
# Croquis: SVG → PNG (resvg vía svg2png-py)
# ─────────────────────────────────────────────────────────────
_FONT_DB_SVG = None


def _font_db_svg():
    """
    FontDatabase cacheado para rasterizar SVG: fuentes del sistema + las
    de `fonts/` del proyecto (para que el texto del croquis tenga glifos
    aunque el sistema no traiga Arial). Import perezoso: solo se necesita
    al convertir un croquis.
    """
    global _FONT_DB_SVG
    if _FONT_DB_SVG is None:
        import svg2png_py

        db = svg2png_py.FontDatabase.system()
        dir_fuentes = Path(__file__).parent / "fonts"
        if dir_fuentes.is_dir():
            try:
                db.load_fonts_dir(str(dir_fuentes))
            except Exception:
                pass
        _FONT_DB_SVG = db
    return _FONT_DB_SVG


def _svg_con_tamano(svg, ancho):
    """
    Fuerza width/height en píxeles en la raíz <svg>. resvg ignora
    width='100%' y renderiza al tamaño del viewBox (unidades), que para
    el croquis es diminuto. Se calcula el alto desde el viewBox para no
    deformar la proporción.
    """
    m = re.search(r'viewBox="[^"]*?\s([\d.]+)\s([\d.]+)"', svg)
    if not m:
        return svg
    vb_ancho, vb_alto = float(m.group(1)), float(m.group(2))
    if vb_ancho <= 0:
        return svg
    alto = max(1, round(ancho * vb_alto / vb_ancho))

    def _fix(m):
        tag = m.group(0)
        tag = re.sub(r'width="[^"]*"', f'width="{ancho}"', tag, count=1)
        tag = re.sub(r'height="[^"]*"', f'height="{alto}"', tag, count=1)
        return tag

    return re.sub(r'<svg\b[^>]*>', _fix, svg, count=1)


def svg_a_png(svg, ancho=1600):
    """
    Convierte un SVG (str) a PNG (bytes) con svg2png-py (resvg), fijando
    `ancho` px para que no salga pixelado al escalarlo en el reporte.
    Devuelve bytes o None si svg2png-py no está disponible o falla.
    """
    if not svg:
        return None
    try:
        import svg2png_py

        return svg2png_py.svg_to_png(_svg_con_tamano(svg, ancho),
                                     _font_db_svg())
    except Exception:
        return None


def _croquis_a_imagenes(ctx):
    """
    ctx['croquis'] = [(título, svg_str)]  →  [(título, png_bytes)].
    Los croquis que no se puedan rasterizar se omiten.
    """
    imagenes = []
    for titulo, svg in ctx.get("croquis") or []:
        png = svg_a_png(svg)
        if png:
            imagenes.append((titulo, png))
    return imagenes


def _aplanar_columnas(columnas):
    """MultiIndex ('Calle | Sentido', 'M1') → 'Calle | Sentido | M1'."""
    salida = []
    for col in columnas:
        if isinstance(col, tuple):
            salida.append(" | ".join(str(p).strip() for p in col
                                     if str(p).strip()))
        else:
            salida.append(str(col).strip())
    return salida


def _df_a_tabla(df, index_nombre="Tipo", encabezados=None):
    """DataFrame → (encabezados, filas) con columnas aplanadas."""
    if encabezados is None:
        encabezados = _aplanar_columnas(df.columns)
    encabezados = [index_nombre] + list(encabezados)
    filas = [[str(idx)] + [_fmt(v) for v in fila] for idx, fila in df.iterrows()]
    return encabezados, filas



def _anchos_proporcionales(encabezados, filas, ancho_util, familia="Helvetica",
                           fuente=7, minima_cm=0.6):
    """
    Anchos de columna según el contenido real (métricas de la fuente),
    normalizados para sumar `ancho_util` exacto.

    Cada columna pide: ancho de la palabra más larga del encabezado
    (negrita) o del dato más largo (regular) + padding de celda (2+2 pt)
    + 2 pt de seguridad. Así ninguna etiqueta se corta a mitad de palabra
    ni se pisan los encabezados. Si sobra página, el excedente se reparte
    en proporción; si falta, se escala (último recurso).
    """
    from reportlab.lib.units import cm
    from reportlab.pdfbase.pdfmetrics import stringWidth

    negrita = "Helvetica-Bold" if familia == "Helvetica" else "DejaVu-Bold"
    extra = 6  # 2 pt × 2 lados de padding + 2 pt de seguridad

    n = len(encabezados)
    if n <= 1:
        return [ancho_util]

    deseado = []
    for j in range(n):
        palabras = [p for parte in str(encabezados[j]).split("|")
                    for p in parte.split()]
        ancho_enc = max((stringWidth(p, negrita, fuente) for p in palabras),
                        default=0)
        ancho_datos = 0
        for fila in filas:
            if j < len(fila):
                ancho_datos = max(ancho_datos,
                                  stringWidth(str(fila[j]), familia, fuente))
        deseado.append(max(ancho_enc, ancho_datos) + extra)

    deseado = [max(d, minima_cm * cm) for d in deseado]

    total = sum(deseado)
    if total <= ancho_util:
        # Sobra página: repartir el excedente en proporción al contenido
        excedente = ancho_util - total
        return [d + excedente * d / total for d in deseado]
    # No entra: escalar (los encabezados más largos pueden envolver más)
    factor = ancho_util / total
    return [d * factor for d in deseado]


def _ficha(ctx):
    """Pares (etiqueta, valor) de los datos manuales del reporte."""
    fecha = ctx.get("fecha")
    if hasattr(fecha, "strftime"):
        fecha_txt = fecha.strftime("%d/%m/%Y")
    else:
        fecha_txt = str(fecha or "")
    return [
        ("Fecha", fecha_txt),
        ("Día", ctx.get("dia", "")),
        ("Ubicación", ctx.get("ubicacion", "")),
        ("Lugar", ctx.get("lugar", "")),
        ("Analista", ctx.get("analista", "")),
    ]


def _metricas_fhp(fhp):
    """Pares (etiqueta, valor) del bloque de hora pico / FHP."""
    return [
        ("Hora pico", f"{fhp['hora_pico_inicio']} – {fhp['hora_pico_fin']}"),
        ("Volumen hora pico", _fmt(fhp["volumen_hp"])),
        ("Bloque máximo de 15 min",
         f"{fhp['bloque_max_inicio']} – {fhp['bloque_max_fin']}"),
        ("Volumen máx. 15 min", _fmt(fhp["volumen_max_15min"])),
        ("FHP", f"{fhp['fhp']:.3f}"),
    ]


# ─────────────────────────────────────────────────────────────
# PDF (reportlab)
# ─────────────────────────────────────────────────────────────
def _registrar_fuentes_pdf():
    """
    Registra DejaVu Sans (embebida en el PDF). Fallback: Helvetica.
    Devuelve el nombre de la familia a usar.
    """
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if "DejaVu" in pdfmetrics.getRegisteredFontNames():
        return "DejaVu"

    for ttf, ttf_bold in _FUENTES:
        if ttf.exists() and ttf_bold.exists():
            try:
                pdfmetrics.registerFont(TTFont("DejaVu", str(ttf)))
                pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(ttf_bold)))
                pdfmetrics.registerFontFamily(
                    "DejaVu", normal="DejaVu", bold="DejaVu-Bold",
                    italic="DejaVu", boldItalic="DejaVu-Bold",
                )
                return "DejaVu"
            except Exception:
                continue
    return "Helvetica"


def _tabla_pdf(datos, col_widths=None, fuente=8, familia="Helvetica",
               encabezado=False, resaltar_total=False):
    """
    Tabla de reportlab con estilos acordes al reporte.

    - Si `encabezado`, la primera fila se convierte en Paragraph para que
      los títulos largos se envuelvan en vez de pisarse entre columnas.
    - Si `resaltar_total`, la última columna (Total) de las filas de datos
      se sombrea en #D6E4F0 (el resto de los datos, en blanco).
    """
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, Table, TableStyle

    datos = [list(fila) for fila in datos]
    if encabezado:
        estilo_enc = ParagraphStyle(
            "EncTabla",
            fontName=("Helvetica-Bold" if familia == "Helvetica"
                      else "DejaVu-Bold"),
            fontSize=fuente,
            leading=fuente + 1.5,
            alignment=1,  # centrado
        )
        datos[0] = [Paragraph(escape(str(c)), estilo_enc) for c in datos[0]]

    tabla = Table(datos, colWidths=col_widths,
                  repeatRows=1 if encabezado else 0)
    estilo = [
        ("FONTNAME", (0, 0), (-1, -1), familia),
        ("FONTSIZE", (0, 0), (-1, -1), fuente),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c0c0c0")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    if encabezado:
        estilo += [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef7")),
        ]
    else:
        estilo += [
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f5f7fa")),
        ]

    # Última columna (Total) sombreada; el resto de los datos, en blanco
    if resaltar_total and len(datos[0]) > 1:
        j = len(datos[0]) - 1
        estilo.append(("BACKGROUND", (j, 1 if encabezado else 0), (j, -1),
                       colors.HexColor("#d6e4f0")))

    tabla.setStyle(TableStyle(estilo))
    return tabla


def _tabla_pdf_multinivel(columnas, filas, enc_nivel2, col_widths,
                          fuente=7, familia="Helvetica"):
    """
    Tabla con encabezado de 2 filas y SPAN (igual que la app de Streamlit):

    - Fila 1: 'Tipo' (ocupa 2 filas) + una celda fusionada por grupo
      ('Calle | Sentido', …; la columna total general, su propio grupo).
    - Fila 2: cada movimiento (M1…, Total); celdas vacías bajo 'Tipo' y
      bajo la columna total general.
    - Datos en blanco; columna 'Total' de cada grupo en #E8EEF7 y la
      Total general en #D6E4F0. repeatRows=2 si la tabla se parte.
    """
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, Table, TableStyle

    negrita = "Helvetica-Bold" if familia == "Helvetica" else "DejaVu-Bold"
    estilo0 = ParagraphStyle("EncN0", fontName=negrita, fontSize=fuente,
                             leading=fuente + 1.5, alignment=1)
    estilo1 = ParagraphStyle("EncN1", fontName=negrita,
                             fontSize=max(fuente - 1, 5),
                             leading=fuente + 1, alignment=1)

    # ── Encabezado: fila 1 (grupos fusionados) y fila 2 (movimientos) ──
    fila0 = [Paragraph("Tipo", estilo0)]
    fila1 = [Paragraph("", estilo1)]
    spans = [("SPAN", (0, 0), (0, 1))]
    totales_grupo, total_general = [], None
    i = 0
    while i < len(columnas):
        col = columnas[i]
        grupo = (str(col[0]).strip() if isinstance(col, tuple)
                 else str(col).strip())
        j = i
        while j + 1 < len(columnas):
            c2 = columnas[j + 1]
            g2 = (str(c2[0]).strip() if isinstance(c2, tuple)
                  else str(c2).strip())
            if g2 != grupo:
                break
            j += 1
        # El grupo va en la celda ancla del span; el resto de las columnas
        # cubiertas se rellena para que la fila tenga una celda por columna
        # (reportlab espera filas del mismo largo que la tabla).
        fila0.append(Paragraph(escape(grupo), estilo0))
        fila0.extend([""] * (j - i))
        if j > i:
            spans.append(("SPAN", (1 + i, 0), (1 + j, 0)))
        for k in range(i, j + 1):
            fila1.append(Paragraph(escape(str(enc_nivel2[1 + k])), estilo1))
        i = j + 1

    # ── Columnas Total (índice de tabla = 1 + posición en `columnas`) ──
    for k, col in enumerate(columnas):
        if not isinstance(col, tuple):
            continue
        n0 = str(col[0]).strip().lower()
        n1 = str(col[1]).strip().lower() if len(col) > 1 else ""
        if n0 == "total":
            total_general = 1 + k
        elif n1 == "total":
            totales_grupo.append(1 + k)

    datos = [fila0, fila1] + [list(f) for f in filas]
    tabla = Table(datos, colWidths=col_widths, repeatRows=2)
    estilo = [
        ("FONTNAME", (0, 2), (-1, -1), familia),
        ("FONTSIZE", (0, 2), (-1, -1), fuente),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c0c0c0")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#e8eef7")),
    ]
    estilo += spans
    for j in totales_grupo:
        estilo.append(("BACKGROUND", (j, 2), (j, -1),
                       colors.HexColor("#e8eef7")))
    if total_general is not None:
        estilo.append(("BACKGROUND", (total_general, 2), (total_general, -1),
                       colors.HexColor("#d6e4f0")))
    tabla.setStyle(TableStyle(estilo))
    return tabla


def _grupo_grafico(titulo):
    """
    Categoría del gráfico por su título, para emparejar SOLO gráficos del
    mismo tipo (p.ej. las 2 tortas juntas, las barras por hora juntas) y no
    mezclar una torta con una barra.
    """
    t = str(titulo).lower()
    if "composición" in t:
        return "tortas"
    if "distribución" in t:
        return "linea"
    if "por hora" in t:
        return "hora"
    if "15 min" in t:
        return "quince"
    if "pesados" in t:
        return "pesados"
    return "otros"


def _historia_imagenes(con_imagenes, ancho_util, negrita):
    """
    Flowables de gráficos para el PDF:

    - De a 2 imágenes por fila (del MISMO grupo), con el título ARRIBA de
      cada una (mismo estilo que los títulos de sección); fluyen las que
      quepan en página.
    - Las imágenes 'solas' (croquis, o más altas que anchas ×1.15) van a
      ancho completo con página propia (KeepTogether + PageBreak).
    - Un grupo con cantidad impar deja su última imagen a ancho completo,
      sin mezclarla con el grupo siguiente.
    """
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (Image, KeepTogether, PageBreak,
                                    Paragraph, Spacer, Table, TableStyle)

    estilo_titulo = ParagraphStyle("TituloGrafico", fontName=negrita,
                                   fontSize=10, leading=12)

    # Clasificar: sola (ancho completo) o emparejable + grupo.
    items = []
    for titulo, png in con_imagenes:
        w, h = _png_dims(png)
        solo = ("croquis" in titulo.lower()) or (h / max(w, 1) > 1.15)
        items.append((titulo, png, w, h, solo, _grupo_grafico(titulo)))

    flujables = []

    def _fila(fila_items, ancho_col):
        """Tabla de 1 o 2 imágenes con sus títulos arriba."""
        titulos, imgs = [], []
        for t, p, iw, ih, *_ in fila_items:
            esc = ancho_col / iw
            titulos.append(Paragraph(escape(_sin_emoji(t)), estilo_titulo))
            imgs.append(Image(io.BytesIO(p), width=ancho_col,
                              height=ih * esc))
        tabla = Table([titulos, imgs], colWidths=[ancho_col] * len(imgs))
        tabla.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        return KeepTogether([tabla])

    def _volcar(pendientes):
        """Empareja los pendientes de a 2 (ya son del mismo grupo)."""
        k = 0
        while k < len(pendientes):
            if k + 1 < len(pendientes):
                ancho_col = (ancho_util - 0.3 * cm) / 2
                flujables.append(_fila(pendientes[k:k + 2], ancho_col))
                k += 2
            else:
                flujables.append(_fila(pendientes[k:k + 1], ancho_util))
                k += 1
            flujables.append(Spacer(1, 6))

    pendientes, grupo_actual = [], None
    for it in items:
        titulo, png, w, h, solo, grupo = it
        if solo:
            _volcar(pendientes)
            pendientes, grupo_actual = [], None
            esc = min(ancho_util / w, 23 * cm / h)
            flujables.append(KeepTogether([
                Paragraph(escape(_sin_emoji(titulo)), estilo_titulo),
                Image(io.BytesIO(png), width=w * esc, height=h * esc),
            ]))
            flujables.append(PageBreak())
            continue
        if pendientes and grupo != grupo_actual:
            _volcar(pendientes)
            pendientes = []
        grupo_actual = grupo
        pendientes.append(it)
    _volcar(pendientes)

    # Sin PageBreak final (evita una página en blanco).
    while flujables and isinstance(flujables[-1], PageBreak):
        flujables.pop()

    return flujables


def generar_pdf(ctx):
    """Genera el reporte PDF (orientación mixta) y devuelve bytes."""
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape, portrait
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (BaseDocTemplate, Frame, NextPageTemplate,
                                    PageBreak, PageTemplate, Paragraph,
                                    Spacer)

    familia = _registrar_fuentes_pdf()
    negrita = "DejaVu-Bold" if familia == "DejaVu" else "Helvetica-Bold"

    estilos = getSampleStyleSheet()
    h1 = ParagraphStyle("RepH1", parent=estilos["Title"], fontName=negrita,
                        fontSize=17, spaceAfter=2)
    sub_st = ParagraphStyle("RepSub", parent=estilos["Normal"],
                            fontName=familia, fontSize=11, alignment=TA_CENTER)
    h2 = ParagraphStyle("RepH2", parent=estilos["Heading2"], fontName=negrita,
                        fontSize=12, spaceBefore=12, spaceAfter=4)

    # ── Documento con dos plantillas: vertical y horizontal ──
    buf = io.BytesIO()
    m = 1.5 * cm
    ancho_v, alto_v = portrait(A4)
    ancho_h, alto_h = landscape(A4)
    doc = BaseDocTemplate(
        buf, pagesize=portrait(A4),
        leftMargin=m, rightMargin=m, topMargin=m, bottomMargin=m,
        title=str(ctx.get("titulo") or "Reporte de censo"),
    )
    doc.addPageTemplates([
        PageTemplate(id="vertical",
                     frames=[Frame(m, m, ancho_v - 2 * m, alto_v - 2 * m,
                                   id="frame_v")],
                     pagesize=portrait(A4)),
        PageTemplate(id="horizontal",
                     frames=[Frame(m, m, ancho_h - 2 * m, alto_h - 2 * m,
                                   id="frame_h")],
                     pagesize=landscape(A4)),
    ])
    ancho_util = ancho_v - 2 * m        # usable vertical (texto, fichas…)
    ancho_util_h = ancho_h - 2 * m      # usable horizontal (tablas anchas)
    plantilla = "vertical"              # plantilla activa
    story = []

    # ── Encabezado ──
    story.append(Paragraph(
        escape(_sin_emoji(ctx.get("titulo") or "Reporte de censo")), h1))
    subtitulo = _sin_emoji(ctx.get("subtitulo"))
    if subtitulo:
        story.append(Paragraph(escape(subtitulo), sub_st))
    story.append(Spacer(1, 10))

    # ── Datos del censo ──
    story.append(Paragraph("Datos del censo", h2))
    ficha = [[escape(k), escape(str(v or "—"))] for k, v in _ficha(ctx)]
    story.append(_tabla_pdf(ficha, col_widths=[3.5 * cm, 12 * cm],
                            fuente=9, familia=familia))

    # ── FHP ──
    story.append(Paragraph("Factor de hora pico (FHP)", h2))
    fhp = [[escape(k), escape(v)] for k, v in _metricas_fhp(ctx["fhp_data"])]
    story.append(_tabla_pdf(fhp, col_widths=[5 * cm, 8 * cm],
                            fuente=9, familia=familia))

    # ── Resumen agrupado — horizontal, encabezado multinivel ──
    resumen_agr = ctx.get("resumen_agr")
    if resumen_agr is not None:
        if plantilla != "horizontal":
            story.append(NextPageTemplate("horizontal"))
            plantilla = "horizontal"
            story.append(PageBreak())
        story.append(Paragraph("Resumen hora pico (agrupado)", h2))
        columnas = list(resumen_agr.columns)
        filas = [[str(idx)] + [_fmt(v) for v in fila]
                 for idx, fila in resumen_agr.iterrows()]
        enc_nivel2 = ["Tipo"] + [
            (str(c[1]).strip() if isinstance(c, tuple) and len(c) > 1 else "")
            for c in columnas
        ]
        story.append(_tabla_pdf_multinivel(
            columnas=columnas, filas=filas, enc_nivel2=enc_nivel2,
            col_widths=_anchos_proporcionales(enc_nivel2, filas,
                                              ancho_util_h,
                                              familia=familia, fuente=7),
            fuente=7, familia=familia,
        ))

    # ── % pesados (si aplica) — vertical ──
    pesados_hp = ctx.get("pesados_hp")
    if pesados_hp is not None:
        if plantilla != "vertical":
            story.append(NextPageTemplate("vertical"))
            plantilla = "vertical"
            story.append(PageBreak())
        story.append(Paragraph("% de vehículos pesados — hora pico", h2))
        enc, filas = _df_a_tabla(pesados_hp, index_nombre="tipo_vehiculo")
        story.append(_tabla_pdf(
            [enc] + filas,
            col_widths=_anchos_proporcionales(enc, filas, ancho_util,
                                              familia=familia, fuente=6.5),
            fuente=6.5, familia=familia, encabezado=True,
            resaltar_total=True,
        ))

    # ── Totales 15 min — horizontal ──
    totales = ctx.get("totales")
    if totales is not None:
        if plantilla != "horizontal":
            story.append(NextPageTemplate("horizontal"))
            plantilla = "horizontal"
            story.append(PageBreak())
        story.append(Paragraph("Totales por bloque de 15 min", h2))
        enc, filas = _df_a_tabla(totales, index_nombre="#")
        fuente_tot = 7 if len(enc) <= 16 else 6
        story.append(_tabla_pdf(
            [enc] + filas,
            col_widths=_anchos_proporcionales(enc, filas, ancho_util_h,
                                              familia=familia,
                                              fuente=fuente_tot),
            fuente=fuente_tot,
            familia=familia, encabezado=True,
            resaltar_total=True,
        ))

    # ── Gráficos — vertical, de a 2 por fila (títulos arriba) ──
    con_imagenes = [(t, png) for t, png in ctx.get("imagenes", []) if png]
    if con_imagenes:
        if plantilla != "vertical":
            story.append(NextPageTemplate("vertical"))
            plantilla = "vertical"
            story.append(PageBreak())
        story.append(Paragraph("Gráficos", h2))
        story += _historia_imagenes(con_imagenes, ancho_util, negrita)

    # ── Croquis — vertical, a ancho completo (una imagen por página) ──
    croquis_pngs = _croquis_a_imagenes(ctx)
    if croquis_pngs:
        if plantilla != "vertical":
            story.append(NextPageTemplate("vertical"))
            plantilla = "vertical"
            story.append(PageBreak())
        story.append(Paragraph("Croquis", h2))
        story += _historia_imagenes(croquis_pngs, ancho_util, negrita)

    doc.build(story)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────
# DOCX (python-docx)
# ─────────────────────────────────────────────────────────────
def _sombrear_celda(celda, color_hex):
    """Fondo de celda (w:shd); python-docx no tiene API nativa para esto."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = celda._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), color_hex.lstrip("#"))
    tc_pr.append(sombra)


def _tabla_docx(doc, encabezados, filas, fuente_pt=8, resaltar_total=False):
    """
    Agrega una tabla estilo 'Table Grid' con fuente Arial controlada.

    Si `resaltar_total`, la última columna (Total) de las filas de datos
    se sombrea en #D6E4F0 (igual que el PDF); el resto, en blanco.
    """
    from docx.shared import Pt

    tabla = doc.add_table(rows=1, cols=len(encabezados))
    tabla.style = "Table Grid"
    for j, txt in enumerate(encabezados):
        tabla.rows[0].cells[j].text = str(txt)
    for fila in filas:
        celdas = tabla.add_row().cells
        for j, txt in enumerate(fila):
            celdas[j].text = str(txt)

    for i, fila in enumerate(tabla.rows):
        for celda in fila.cells:
            for parrafo in celda.paragraphs:
                parrafo.paragraph_format.space_before = Pt(0)
                parrafo.paragraph_format.space_after = Pt(0)
                for run in parrafo.runs:
                    run.font.size = Pt(fuente_pt)
                    run.font.name = "Arial"
                    if i == 0:
                        run.bold = True

    # Encabezado con el mismo fondo que el PDF
    for celda in tabla.rows[0].cells:
        _sombrear_celda(celda, "#e8eef7")

    # Última columna (Total) sombreada; el resto de los datos, en blanco
    if resaltar_total and len(encabezados) > 1:
        j = len(encabezados) - 1
        for fila in tabla.rows[1:]:
            _sombrear_celda(fila.cells[j], "#d6e4f0")
    return tabla


def _tabla_docx_multinivel(doc, columnas, filas, enc_nivel2, fuente_pt=7):
    """
    Tabla con encabezado de 2 filas fusionadas (igual que el PDF):

    - Fila 1: 'Tipo' (fusión vertical) + una celda por grupo fusionada.
    - Fila 2: cada movimiento; fila vacía bajo 'Tipo' y bajo la Total
      general.
    - Datos en blanco; 'Total' de cada grupo en #E8EEF7 y Total general
      en #D6E4F0.
    """
    from docx.shared import Pt

    ncols = len(enc_nivel2)
    tabla = doc.add_table(rows=2, cols=ncols)
    tabla.style = "Table Grid"

    # Fila 2: movimientos (la columna 0 queda bajo la fusión de "Tipo")
    for j in range(1, ncols):
        tabla.rows[1].cells[j].text = str(enc_nivel2[j])

    # Fila 1: grupos fusionados
    i = 0
    while i < len(columnas):
        col = columnas[i]
        grupo = (str(col[0]).strip() if isinstance(col, tuple)
                 else str(col).strip())
        j = i
        while j + 1 < len(columnas):
            c2 = columnas[j + 1]
            g2 = (str(c2[0]).strip() if isinstance(c2, tuple)
                  else str(c2).strip())
            if g2 != grupo:
                break
            j += 1
        celda = tabla.rows[0].cells[1 + i]
        if j > i:
            celda = celda.merge(tabla.rows[0].cells[1 + j])
        celda.text = grupo
        i = j + 1

    # Fusión vertical de "Tipo"
    celda_tipo = tabla.rows[0].cells[0].merge(tabla.rows[1].cells[0])
    celda_tipo.text = "Tipo"

    # Filas de datos
    for fila in filas:
        celdas = tabla.add_row().cells
        for j, txt in enumerate(fila):
            celdas[j].text = str(txt)

    # Fuentes: Arial; encabezados en negrita
    for i_fila, fila in enumerate(tabla.rows):
        for celda in fila.cells:
            for parrafo in celda.paragraphs:
                parrafo.paragraph_format.space_before = Pt(0)
                parrafo.paragraph_format.space_after = Pt(0)
                for run in parrafo.runs:
                    run.font.size = Pt(fuente_pt)
                    run.font.name = "Arial"
                    if i_fila < 2:
                        run.bold = True

    # Encabezados con fondo (dedupe de celdas fusionadas)
    vistos = set()
    for fila in tabla.rows[:2]:
        for celda in fila.cells:
            if id(celda._tc) in vistos:
                continue
            vistos.add(id(celda._tc))
            _sombrear_celda(celda, "#e8eef7")

    # Banding: sólo columnas Total (índice de tabla = 1 + posición)
    for k, col in enumerate(columnas):
        if not isinstance(col, tuple):
            continue
        n0 = str(col[0]).strip().lower()
        n1 = str(col[1]).strip().lower() if len(col) > 1 else ""
        if n0 == "total":
            color = "#d6e4f0"
        elif n1 == "total":
            color = "#e8eef7"
        else:
            continue
        for fila in tabla.rows[2:]:
            _sombrear_celda(fila.cells[1 + k], color)
    return tabla


def _pares_docx(doc, pares):
    """Tabla de 2 columnas (concepto, valor) a partir de pares."""
    _tabla_docx(doc, ["Concepto", "Valor"],
                [[str(k), str(v or "—")] for k, v in pares],
                fuente_pt=9)


def generar_docx(ctx):
    """Genera el reporte DOCX y devuelve bytes."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    doc = Document()

    # ── Horizontal + Arial en los estilos base ──
    seccion = doc.sections[0]
    seccion.orientation = WD_ORIENT.LANDSCAPE
    seccion.page_width, seccion.page_height = (seccion.page_height,
                                               seccion.page_width)
    for nombre in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        try:
            doc.styles[nombre].font.name = "Arial"
        except KeyError:
            pass
    doc.styles["Normal"].font.size = Pt(10)

    # ── Encabezado ──
    doc.add_heading(str(ctx.get("titulo") or "Reporte de censo"), level=0)
    subtitulo = str(ctx.get("subtitulo") or "")
    if subtitulo:
        parrafo = doc.add_paragraph(subtitulo)
        parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ── Datos del censo ──
    doc.add_heading("Datos del censo", level=1)
    _pares_docx(doc, _ficha(ctx))

    # ── FHP ──
    doc.add_heading("Factor de hora pico (FHP)", level=1)
    _pares_docx(doc, _metricas_fhp(ctx["fhp_data"]))

    # ── Resumen agrupado (encabezado de 2 filas) ──
    resumen_agr = ctx.get("resumen_agr")
    if resumen_agr is not None:
        doc.add_heading("Resumen hora pico (agrupado)", level=1)
        columnas = list(resumen_agr.columns)
        filas = [[str(idx)] + [_fmt(v) for v in fila]
                 for idx, fila in resumen_agr.iterrows()]
        enc_nivel2 = ["Tipo"] + [
            (str(c[1]).strip() if isinstance(c, tuple) and len(c) > 1 else "")
            for c in columnas
        ]
        _tabla_docx_multinivel(doc, columnas, filas, enc_nivel2)

    # ── % pesados ──
    pesados_hp = ctx.get("pesados_hp")
    if pesados_hp is not None:
        doc.add_heading("% de vehículos pesados — hora pico", level=1)
        enc, filas = _df_a_tabla(pesados_hp, index_nombre="tipo_vehiculo")
        _tabla_docx(doc, enc, filas, fuente_pt=7, resaltar_total=True)

    # ── Totales 15 min ──
    totales = ctx.get("totales")
    if totales is not None:
        doc.add_heading("Totales por bloque de 15 min", level=1)
        enc, filas = _df_a_tabla(totales, index_nombre="#")
        _tabla_docx(doc, enc, filas, fuente_pt=7, resaltar_total=True)

    # ── Gráficos ──
    con_imagenes = [(t, png) for t, png in ctx.get("imagenes", []) if png]
    if con_imagenes:
        doc.add_heading("Gráficos", level=1)
    ancho_max_cm, alto_max_cm = 24.0, 14.5
    for titulo, png in con_imagenes:
        doc.add_heading(str(titulo), level=2)
        w_px, h_px = _png_dims(png)
        escala = min(ancho_max_cm / w_px, alto_max_cm / h_px)
        doc.add_picture(io.BytesIO(png), width=Cm(w_px * escala))

    # ── Croquis ──
    croquis_pngs = _croquis_a_imagenes(ctx)
    if croquis_pngs:
        doc.add_heading("Croquis", level=1)
    for titulo, png in croquis_pngs:
        doc.add_heading(str(titulo), level=2)
        w_px, h_px = _png_dims(png)
        escala = min(ancho_max_cm / w_px, alto_max_cm / h_px)
        doc.add_picture(io.BytesIO(png), width=Cm(w_px * escala))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def gen_pdf(imgs, output_path, title, subtitle, date, croquis=None):
    """Wrapper function for test compatibility. Genera PDF y guarda en archivo."""
    ctx = {
        "titulo": title,
        "subtitulo": subtitle,
        "fecha": date,
        "imagenes": imgs,
        "croquis": croquis or [],
        "fhp_data": {
            "fhp": 0.813,
            "hora_pico_inicio": "07:30",
            "hora_pico_fin": "08:30",
            "volumen_hp": 1000,
            "volumen_max_15min": 250,
            "bloque_max_inicio": "07:45",
            "bloque_max_fin": "08:00"
        },
        "resumen_agr": None,
        "resumen_plano": None,
        "pesados_hp": None,
        "totales": None,
        "excluidos": [],
        "ubicacion": "",
        "lugar": "",
        "analista": "",
        "dia": ""
    }
    pdf_bytes = generar_pdf(ctx)
    with open(output_path, "wb") as f:
        f.write(pdf_bytes)
    return output_path


def gen_docx(imgs, output_path, title, subtitle, date, croquis=None):
    """Wrapper function for test compatibility. Genera DOCX y guarda en archivo."""
    ctx = {
        "titulo": title,
        "subtitulo": subtitle,
        "fecha": date,
        "imagenes": imgs,
        "croquis": croquis or [],
        "fhp_data": {
            "fhp": 0.813,
            "hora_pico_inicio": "07:30",
            "hora_pico_fin": "08:30",
            "volumen_hp": 1000,
            "volumen_max_15min": 250,
            "bloque_max_inicio": "07:45",
            "bloque_max_fin": "08:00"
        },
        "resumen_agr": None,
        "resumen_plano": None,
        "pesados_hp": None,
        "totales": None,
        "excluidos": [],
        "ubicacion": "",
        "lugar": "",
        "analista": "",
        "dia": ""
    }
    docx_bytes = generar_docx(ctx)
    with open(output_path, "wb") as f:
        f.write(docx_bytes)
    return output_path



